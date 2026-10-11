from __future__ import annotations

import logging
import re
from pathlib import Path

logger = logging.getLogger(__name__)


class DocumentLoader:
    """
    Direct document reader for CV files.
    Supports native digital text extraction from PDF, DOCX, and TXT files.
    """

    SUPPORTED_EXTENSIONS = {
        ".pdf": "pdf",
        ".docx": "docx",
        ".doc": "doc",
        ".docs": "doc",
        ".txt": "text",
        ".md": "text",
    }

    def load_text(self, file_path: str) -> tuple[str, str, list[str]]:
        """
        Loads and extracts text and embedded hyperlinks from supported digital CV documents.
        Returns: (extracted_text, document_format, extracted_links)
        All extracted links are returned locally per call, avoiding shared mutable state.
        """
        path = Path(file_path)
        if not path.exists():
            raise FileNotFoundError(f"CV file not found: {file_path}")

        ext = path.suffix.lower()
        doc_type = self.SUPPORTED_EXTENSIONS.get(ext)

        if not doc_type:
            raise ValueError(
                f"Unsupported document format: '{ext}'. "
                f"Supported formats are: {', '.join(sorted(self.SUPPORTED_EXTENSIONS.keys()))}"
            )

        if doc_type == "pdf":
            text, links = self._extract_pdf(path)
            return text, "pdf", links
        elif doc_type == "docx":
            text, links = self._extract_docx(path)
            return text, "docx", links
        elif doc_type == "doc":
            text, links = self._extract_doc(path)
            return text, "doc", links
        elif doc_type == "text":
            text, links = self._extract_plain_text(path)
            return text, "text", links
        else:
            raise ValueError(f"No extractor available for format: {ext}")

    def _extract_pdf(self, path: Path) -> tuple[str, list[str]]:
        """Extract readable PDF text while returning hyperlink annotations separately."""
        text_pages = []
        extracted_links: list[str] = []
        try:
            import pypdf
            reader = pypdf.PdfReader(str(path))
            for page in reader.pages:
                # 1. Collect link annotations. Project association receives these separately.
                if hasattr(page, "annotations") and page.annotations:
                    for annot_ref in page.annotations:
                        try:
                            annot = annot_ref.get_object() if hasattr(annot_ref, "get_object") else annot_ref
                            if not annot or annot.get("/Subtype") != "/Link":
                                continue
                            action = annot.get("/A")
                            uri = None
                            if action and hasattr(action, "get"):
                                uri = action.get("/URI")
                            elif annot and hasattr(annot, "get") and "/URI" in annot:
                                uri = annot.get("/URI")

                            if uri:
                                uri_str = str(uri).strip()
                                if uri_str:
                                    if uri_str not in extracted_links:
                                        extracted_links.append(uri_str)
                        except Exception:
                            continue

                # Layout extraction is substantially more reliable for glyph-positioned text.
                # Do not assume it is usable for every PDF: retain standard extraction as a
                # generic fallback when layout loses a material amount of textual content.
                try:
                    layout_text = (page.extract_text(extraction_mode="layout") or "").strip()
                except Exception as layout_error:
                    logger.warning("PDF layout extraction failed; using standard extraction: %s", type(layout_error).__name__)
                    layout_text = ""

                try:
                    standard_text = (page.extract_text() or "").strip()
                except Exception as standard_error:
                    logger.warning("PDF standard extraction failed: %s", type(standard_error).__name__)
                    standard_text = ""

                page_text = self._select_pdf_text(layout_text, standard_text)

                if page_text:
                    text_pages.append(page_text)

        except Exception as e:
            logger.error(f"pypdf extraction failed for {path}: {e}")
            raise ValueError(f"Failed to read PDF document '{path.name}': {e!s}")

        full_text = "\n".join(text_pages).strip()
        if not full_text:
            raise ValueError(
                f"No readable text could be extracted from '{path.name}'. "
                "Please ensure the document contains digital selectable text (scanned image PDFs are not supported)."
            )
        return full_text, extracted_links

    @staticmethod
    def _select_pdf_text(layout_text: str, standard_text: str) -> str:
        """Prefer layout text only when it retains enough readable content."""
        if not layout_text:
            return standard_text
        if not standard_text:
            return layout_text

        def visible_characters(value: str) -> int:
            return len(re.sub(r"\s+", "", value))

        layout_visible = visible_characters(layout_text)
        standard_visible = visible_characters(standard_text)
        if standard_visible and layout_visible < standard_visible * 0.8:
            logger.warning(
                "PDF layout extraction lost material text; using standard extraction instead: layout_chars=%d standard_chars=%d",
                layout_visible,
                standard_visible,
            )
            return standard_text
        return layout_text

    def _extract_docx(self, path: Path) -> tuple[str, list[str]]:
        """Extracts paragraphs, hyperlinks, and table cells from DOCX files."""
        lines: list[str] = []
        extracted_links: list[str] = []
        try:
            import docx
            from docx.oxml.ns import qn
            doc = docx.Document(str(path))
            for p in doc.paragraphs:
                p_text = ""
                for child in p._p:
                    if child.tag.endswith("hyperlink"):
                        r_id = child.get(qn("r:id"))
                        # In docx oxml, text is strictly in <w:t> nodes.
                        # Do NOT use child.iter() without filtering for w:t, because parent oxml elements
                        # return the same text as their children, causing duplicated/tripled words.
                        link_text = "".join(t.text for t in child.iter(qn("w:t")) if t.text)
                        target = None
                        if r_id and r_id in p.part.rels:
                            target = p.part.rels[r_id].target_ref
                            if target and target not in extracted_links:
                                extracted_links.append(target)
                        if target and target not in link_text:
                            p_text += f" {link_text} ({target}) "
                        else:
                            p_text += link_text
                    elif child.tag.endswith("r"):
                        # In docx oxml, text within a run <w:r> is strictly in <w:t> nodes.
                        run_text = "".join(t.text for t in child.iter(qn("w:t")) if t.text)
                        p_text += run_text
                clean_line = p_text.strip() or p.text.strip()
                if clean_line:
                    lines.append(clean_line)

            for table in doc.tables:
                for row in table.rows:
                    row_text = " | ".join(cell.text.strip() for cell in row.cells if cell.text.strip())
                    if row_text:
                        lines.append(row_text)
            extracted = "\n\n".join(lines).strip()
            if not extracted:
                raise ValueError(
                    f"No readable text could be extracted from DOCX file '{path.name}'."
                )
            return extracted, extracted_links
        except ImportError:
            # Fallback direct zip XML extraction if python-docx isn't installed
            import xml.etree.ElementTree as ET
            import zipfile
            with zipfile.ZipFile(str(path)) as z:
                xml_content = z.read("word/document.xml")
            tree = ET.fromstring(xml_content)
            namespaces = {"w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main"}
            text_nodes = tree.findall(".//w:t", namespaces)
            extracted = " ".join(node.text for node in text_nodes if node.text).strip()
            if not extracted:
                raise ValueError(
                    f"No readable text could be extracted from DOCX file '{path.name}'."
                )
            return extracted, []
        except Exception as e:
            logger.error(f"DOCX extraction error for {path}: {e}")
            raise ValueError(f"Failed to extract text from DOCX file '{path.name}': {e!s}")

    def _extract_plain_text(self, path: Path) -> tuple[str, list[str]]:
        """Reads plain text files with UTF-8 encoding."""
        with open(str(path), "r", encoding="utf-8", errors="replace") as f:
            return f.read(), []

    def _extract_doc(self, path: Path) -> tuple[str, list[str]]:
        """
        Extracts text and embedded links from Word .doc and .docs files.
        Handles:
        1. Renamed DOCX files (Office Open XML zip).
        2. Binary OLE Compound Files (Word 97-2003) via FIB piece table parsing with fallback.
        3. RTF files saved as .doc / .docs.
        """
        import zipfile

        # 1. Check if the file is actually a DOCX archive saved with .doc / .docs extension
        if zipfile.is_zipfile(str(path)):
            try:
                return self._extract_docx(path)
            except Exception as e:
                logger.debug("Attempt to extract %s as zip/docx failed: %s", path.name, e)

        extracted_links: list[str] = []
        extracted_text = ""

        # 2. Try OLE binary Word extraction (Word 97-2003 binary format)
        try:
            import olefile
            if olefile.isOleFile(str(path)):
                with olefile.OleFileIO(str(path)) as ole:
                    if ole.exists("WordDocument"):
                        word_stream = ole.openstream("WordDocument").read()
                        table_name = None
                        if len(word_stream) >= 512:
                            import struct
                            flags = struct.unpack("<H", word_stream[0x000A:0x000C])[0]
                            table_name = "1Table" if (flags & 0x0200) else "0Table"

                        table_stream = b""
                        if table_name and ole.exists(table_name):
                            table_stream = ole.openstream(table_name).read()

                        if table_stream:
                            extracted_text = self._parse_doc_piece_table(word_stream, table_stream)

                        if not extracted_text or len(extracted_text.strip()) < 10:
                            extracted_text = self._extract_doc_stream_fallback(word_stream)
        except Exception as ole_err:
            logger.warning("OLE .doc extraction attempt failed for %s: %s", path.name, ole_err)

        # 3. Fallback for RTF or plain text files saved as .doc / .docs
        if not extracted_text or not extracted_text.strip():
            try:
                raw_bytes = path.read_bytes()
                if raw_bytes.startswith(b"{\\rtf"):
                    extracted_text = self._extract_rtf_text(raw_bytes)
                else:
                    extracted_text = raw_bytes.decode("utf-8", errors="ignore").strip()
            except Exception as read_err:
                logger.debug("Fallback byte decoding failed for %s: %s", path.name, read_err)

        if not extracted_text or not extracted_text.strip():
            raise ValueError(
                f"No readable text could be extracted from '{path.name}'. "
                "Please ensure the document contains digital selectable text."
            )

        extracted_text = extracted_text.replace("\r\n", "\n").replace("\r", "\n")
        extracted_text = re.sub(r"\n{3,}", "\n\n", extracted_text).strip()

        # Extract embedded URLs
        urls = re.findall(r"https?://[^\s<>\"'{}|\\^`]+", extracted_text)
        for u in urls:
            u_clean = u.rstrip(".,;)]")
            if u_clean and u_clean not in extracted_links:
                extracted_links.append(u_clean)

        return extracted_text, extracted_links

    @staticmethod
    def _parse_doc_piece_table(word_stream: bytes, table_stream: bytes) -> str:
        """Parses the Word FIB piece table (clx) to reconstruct text."""
        import struct
        try:
            fcClx = struct.unpack("<I", word_stream[0x01A2:0x01A6])[0]
            lcbClx = struct.unpack("<I", word_stream[0x01A6:0x01AA])[0]
            pos = fcClx
            end = min(fcClx + lcbClx, len(table_stream))
            pieces: list[str] = []
            while pos < end:
                clxt = table_stream[pos]
                pos += 1
                if clxt == 1:
                    if pos + 2 > end:
                        break
                    cbGrpprl = struct.unpack("<H", table_stream[pos:pos + 2])[0]
                    pos += 2 + cbGrpprl
                elif clxt == 2:
                    if pos + 4 > end:
                        break
                    lcb = struct.unpack("<I", table_stream[pos:pos + 4])[0]
                    pos += 4
                    if lcb < 4:
                        break
                    num_pieces = (lcb - 4) // 12
                    cpr_pos = pos
                    pcd_pos = pos + (num_pieces + 1) * 4
                    for i in range(num_pieces):
                        cp_start = struct.unpack("<I", table_stream[cpr_pos + i * 4 : cpr_pos + (i + 1) * 4])[0]
                        cp_end = struct.unpack("<I", table_stream[cpr_pos + (i + 1) * 4 : cpr_pos + (i + 2) * 4])[0]
                        if pcd_pos + i * 8 + 6 > len(table_stream):
                            break
                        fc_desc = struct.unpack("<I", table_stream[pcd_pos + i * 8 + 2 : pcd_pos + i * 8 + 6])[0]
                        is_compressed = bool(fc_desc & 0x40000000)
                        fc = fc_desc & 0xBFFFFFFF
                        length = cp_end - cp_start
                        if length <= 0 or length > 5000000:
                            continue
                        if is_compressed:
                            actual_fc = fc // 2
                            if actual_fc + length <= len(word_stream):
                                raw = word_stream[actual_fc : actual_fc + length]
                                pieces.append(raw.decode("cp1252", errors="replace"))
                        else:
                            if fc + length * 2 <= len(word_stream):
                                raw = word_stream[fc : fc + length * 2]
                                pieces.append(raw.decode("utf-16le", errors="replace"))
                    break
                else:
                    break
            text = "".join(pieces)
            return text.replace("\x07", "\t").replace("\x0b", "\n").replace("\x0c", "\n")
        except Exception as e:
            logger.debug("Error in _parse_doc_piece_table: %s", e)
            return ""

    @staticmethod
    def _extract_doc_stream_fallback(stream: bytes) -> str:
        """Extract readable UTF-16LE and ASCII text runs from raw binary stream."""
        runs: list[str] = []
        utf16_runs = re.findall(rb"(?:[\x20-\x7e\n\r\t]\x00){4,}", stream)
        for r in utf16_runs:
            try:
                decoded = r.decode("utf-16le", errors="ignore").strip()
                if len(decoded) > 3:
                    runs.append(decoded)
            except Exception:
                pass

        if runs:
            return "\n".join(runs)

        ascii_runs = re.findall(rb"[\x20-\x7e\n\r\t]{4,}", stream)
        for r in ascii_runs:
            try:
                decoded = r.decode("cp1252", errors="ignore").strip()
                if len(decoded) > 3 and not decoded.startswith(("{\\", "<!", "<?")):
                    runs.append(decoded)
            except Exception:
                pass

        return "\n".join(runs)

    @staticmethod
    def _extract_rtf_text(rtf_bytes: bytes) -> str:
        """Simple text stripper for RTF formatted documents."""
        text = rtf_bytes.decode("cp1252", errors="replace")
        text = re.sub(r"\\par[d]?", "\n", text)
        text = re.sub(r"\\tab", "\t", text)
        text = re.sub(r"\\[a-zA-Z]+(-?\d+)? ?", "", text)
        text = re.sub(r"[{}]", "", text)
        return text.strip()
